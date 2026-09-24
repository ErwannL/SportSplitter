"""Import de l'emploi du temps Excel et export des plannings générés."""

from __future__ import annotations

import io
import re
from collections import OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .messages import render
from .schemas import (MODE_PERIODS, PERIODS, SEGMENTS, Cell, Entry,
                      Solution, Timetable, TimeRow, Workspace)


class TimetableError(ValueError):
    pass


CLOSED_WORDS = {"x", "fermé", "ferme", "fermée", "closed", "-", "/"}
_TIME_SPLIT = re.compile(r"\s*(?:-|–|—|à|\n|/)\s*")
_GROUPS = [
    re.compile(r"^(?P<name>.+?)\s*[x×*]\s*(?P<n>\d+)$", re.I),
    re.compile(r"^(?P<n>\d+)\s*[x×*]\s*(?P<name>.+)$", re.I),
    re.compile(r"^(?P<name>.+?)\s*\((?P<n>\d+)\)$"),
]


def _fmt_time(v) -> str:
    return f"{v.hour}h" + (f"{v.minute:02d}" if v.minute else "")


def parse_time_label(value) -> TimeRow:
    if hasattr(value, "hour"):
        t = _fmt_time(value)
        return TimeRow(label=t, start=t, end="")
    text = str(value).strip()
    parts = [p for p in _TIME_SPLIT.split(text) if p]
    start = parts[0] if parts else text
    end = parts[-1] if len(parts) > 1 else ""
    label = f"{start} – {end}" if end else start
    return TimeRow(label=label, start=start, end=end)


_CLOCK = re.compile(r"^(\d{1,2})\s*(?:[hH:]\s*(\d{1,2})?)?(?::\d{2})?$")


def to_minutes(text: str) -> int | None:
    """« 8h », « 8h30 », « 08:30 », « 8 » -> minutes depuis minuit."""
    m = _CLOCK.match(text.strip())
    if not m:
        return None
    h, mn = int(m.group(1)), int(m.group(2) or 0)
    return h * 60 + mn if h < 24 and mn < 60 else None


def fill_durations(rows: list[TimeRow]) -> None:
    """Complète la fin manquante par le début de la ligne suivante et calcule la durée."""
    for i, row in enumerate(rows):
        if not row.end and i + 1 < len(rows):
            row.end = rows[i + 1].start
        start, end = to_minutes(row.start), to_minutes(row.end)
        row.minutes = end - start if start is not None and end is not None and end > start else 0
        if row.end and row.label == row.start:
            row.label = f"{row.start} – {row.end}"


def parse_entries(text: str) -> list[Entry]:
    counts: OrderedDict[str, int] = OrderedDict()
    display: dict[str, str] = {}
    for raw in re.split(r"[,;\n+]", text):
        token = raw.strip()
        if not token:
            continue
        name, n = token, 1
        for rx in _GROUPS:
            m = rx.match(token)
            if m:
                name, n = m.group("name").strip(), int(m.group("n"))
                break
        key = name.lower()
        display.setdefault(key, name)
        counts[key] = counts.get(key, 0) + max(n, 1)
    return [Entry(level=display[k], groups=v) for k, v in counts.items()]


def parse_timetable(data: bytes, file_name: str = "") -> Timetable:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise TimetableError("Fichier Excel illisible.") from exc
    ws = wb.worksheets[0]

    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if sum(1 for v in vals[1:] if v not in (None, "")) >= 1 and (
            vals[0] is None or "heure" in str(vals[0]).lower() or r == 1
        ):
            if any(isinstance(v, str) and v.strip() for v in vals[1:]):
                header_row = r
                break
    if header_row is None:
        raise TimetableError("Impossible de trouver la ligne d'en-tête (Heures, Lundi, Mardi…).")

    day_cols: list[tuple[int, str]] = []
    for c in range(2, ws.max_column + 1):
        v = ws.cell(header_row, c).value
        if v not in (None, "") and str(v).strip():
            day_cols.append((c, str(v).strip()))

    row_idx: list[int] = []
    rows: list[TimeRow] = []
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or not str(v).strip():
            if rows:
                break  # fin du tableau
            continue
        row_idx.append(r)
        rows.append(parse_time_label(v))
    if not rows:
        raise TimetableError("Aucun créneau horaire trouvé dans la première colonne.")

    fill_durations(rows)
    merged: dict[tuple[int, int], tuple[int, int, int]] = {}  # (r,c) -> (top, left, bottom)
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = (rng.min_row, rng.min_col, rng.max_row)

    row_pos = {r: i for i, r in enumerate(row_idx)}
    cells: list[Cell] = []
    for d, (c, _) in enumerate(day_cols):
        covered: set[int] = set()
        for i, r in enumerate(row_idx):
            if i in covered:
                continue
            span, is_merged = 1, False
            if (r, c) in merged:
                top, left, bottom = merged[(r, c)]
                # lignes contiguës : les suites de fusion sont déjà dans ``covered``
                if (top, left) != (r, c) and top in row_pos and top != r:  # pragma: no cover
                    continue
                is_merged = True
                span = max(1, sum(1 for rr in row_idx if r <= rr <= bottom))
                for k in range(i + 1, i + span):
                    covered.add(k)
                value = ws.cell(top, left).value
            else:
                value = ws.cell(r, c).value
            text = "" if value is None else str(value).strip()
            closed = text.lower() in CLOSED_WORDS or (is_merged and span > 1 and not text)
            entries = [] if closed else parse_entries(text)
            cells.append(Cell(day=d, row=i, rowSpan=span, closed=closed, entries=entries))

    return Timetable(days=[n for _, n in day_cols], rows=rows, cells=cells, fileName=file_name)


# ---------------------------------------------------------------- export

LABELS = {
    "fr": {"hours": "Heures", "level": "Niveau", "mode": "Organisation", "broken": "Règles non respectées",
           "summary": "Récapitulatif", "solution": "Solution", "week": "Semaine", "trimestre": "Trimestre", "semestre": "Semestre",
           "Q1": "Sept - Nov", "Q2": "Déc - Janv", "Q3": "Févr - Mars", "Q4": "Avr - Juin"},
    "en": {"hours": "Hours", "level": "Level", "mode": "Organisation", "broken": "Broken rules",
           "summary": "Summary", "solution": "Solution", "week": "Week", "trimestre": "Term", "semestre": "Semester",
           "Q1": "Sep - Nov", "Q2": "Dec - Jan", "Q3": "Feb - Mar", "Q4": "Apr - Jun"},
}

THIN = Side(style="thin", color="94A3B8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1E293B")
HEAD_FONT = Font(bold=True, color="FFFFFF")
CLOSED_FILL = PatternFill("solid", fgColor="E2E8F0")
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _hex(color: str) -> str:
    c = color.lstrip("#")
    return ("".join(ch * 2 for ch in c) if len(c) == 3 else c)[:6].upper() or "FFFFFF"


def _text_color(color: str) -> str:
    c = _hex(color)
    r, g, b = int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
    return "000000" if (0.299 * r + 0.587 * g + 0.114 * b) > 150 else "FFFFFF"


def cell_contents(ws: Workspace, sol: Solution, segment: str, week: int = 0) -> dict[tuple[int, int], list]:
    """(jour, ligne) -> [(texte, couleur du lieu)] pour un segment de l'année et une semaine du cycle commun."""
    levels = {lv.id: lv for lv in ws.levels}
    sports = {s.id: s for s in ws.sports}
    places = {p.id: p for p in ws.places}
    out: dict[tuple[int, int], list[tuple[str, str]]] = {}
    for a in sol.assignments:
        lv = levels[a.level_id]
        if segment not in PERIODS[a.period] or week % len(lv.cycle) != a.week:
            continue
        for pl in a.placements:
            place = places[pl.place_id]
            grp = f" ×{pl.groups}" if pl.groups > 1 else ""
            txt = f"{lv.name}{grp} · {sports[a.sport_id].name} · {place.name}"
            for r in range(a.row, a.row + a.span):
                out.setdefault((a.day, r), []).append((txt, place.color))
    return out


def _write_grid(sheet, ws: Workspace, sol: Solution, segment: str, top: int, lang: str = "fr", week: int = 0) -> int:
    tt = ws.timetable
    L = LABELS[lang]
    title = L[segment] if sol.weeks == 1 else f"{L[segment]} · {L['week']} {chr(65 + week)}"
    sheet.cell(top, 1, title).font = Font(bold=True, size=13)
    top += 1
    sheet.cell(top, 1, L["hours"])
    for d, day in enumerate(tt.days):
        sheet.cell(top, d + 2, day)
    for c in range(1, len(tt.days) + 2):
        cell = sheet.cell(top, c)
        cell.fill, cell.font, cell.border, cell.alignment = HEAD_FILL, HEAD_FONT, BORDER, WRAP
    for i, row in enumerate(tt.rows):
        c = sheet.cell(top + 1 + i, 1, row.label)
        c.border, c.alignment, c.font = BORDER, WRAP, Font(bold=True)
        sheet.row_dimensions[top + 1 + i].height = 45
    contents = cell_contents(ws, sol, segment, week)
    for cell in tt.cells:
        r, col = top + 1 + cell.row, cell.day + 2
        if cell.row_span > 1:
            sheet.merge_cells(start_row=r, start_column=col, end_row=r + cell.row_span - 1, end_column=col)
        target = sheet.cell(r, col)
        target.border, target.alignment = BORDER, WRAP
        items = contents.get((cell.day, cell.row), [])
        if cell.closed:
            target.fill = CLOSED_FILL
        elif items:
            target.value = "\n".join(t for t, _ in items)
            target.fill = PatternFill("solid", fgColor=_hex(items[0][1]))
            target.font = Font(color=_text_color(items[0][1]), size=9)
    return top + len(tt.rows) + 3


def _write_summary(sheet, ws: Workspace, sol: Solution, top: int = 1, lang: str = "fr") -> int:
    L = LABELS[lang]
    sports = {s.id: s.name for s in ws.sports}
    sheet.cell(top, 1, L["level"])
    sheet.cell(top, 2, L["mode"])
    for i, p in enumerate(["T1 / S1", "T2 / S2", "T3"]):
        sheet.cell(top, 3 + i, p)
    for c in range(1, 6):
        cell = sheet.cell(top, c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
    r = top + 1
    for lv in ws.levels:
        if lv.id not in sol.plan:
            continue
        sheet.cell(r, 1, lv.name).border = BORDER
        sheet.cell(r, 2, L[lv.mode]).border = BORDER
        for i, p in enumerate(MODE_PERIODS[lv.mode]):
            sheet.cell(r, 3 + i, sports.get(sol.plan[lv.id].get(p, ""), "")).border = BORDER
        r += 1
    if sol.violations:
        r += 1
        sheet.cell(r, 1, L["broken"]).font = Font(bold=True, color="B91C1C")
        for v in sol.violations:
            r += 1
            sheet.cell(r, 1, render(v.rule, v.params, lang) if v.params else v.message)
    return r + 2


def _widths(sheet, n_days: int):
    sheet.column_dimensions["A"].width = 14
    for d in range(n_days + 4):
        sheet.column_dimensions[get_column_letter(d + 2)].width = 30


def export_solutions(ws: Workspace, sols: list[Solution], lang: str = "fr") -> bytes:
    if ws.timetable is None:
        raise TimetableError("Aucun emploi du temps.")
    L = LABELS[lang]
    wb = Workbook()
    wb.remove(wb.active)
    n = len(ws.timetable.days)
    if len(sols) == 1:
        sol = sols[0]
        s = wb.create_sheet(L["summary"])
        _write_summary(s, ws, sol, lang=lang)
        _widths(s, 3)
        for seg in SEGMENTS:
            sh = wb.create_sheet(L[seg])
            top = 1
            for week in range(sol.weeks):
                top = _write_grid(sh, ws, sol, seg, top, lang, week)
            _widths(sh, n)
    else:
        for sol in sols:
            sh = wb.create_sheet(f"{L['solution']} {sol.index + 1}")
            top = _write_summary(sh, ws, sol, lang=lang)
            for seg in SEGMENTS:
                for week in range(sol.weeks):
                    top = _write_grid(sh, ws, sol, seg, top, lang, week)
            _widths(sh, n)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_workbook() -> bytes:
    """Grille vide au pas de 30 min : 8h-18h, pause de midi et mercredi après-midi fermés."""
    wb = Workbook()
    sh = wb.active
    sh.title = "Emploi du temps"
    days = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi"]
    sh.append(["Heures", *days])
    fmt = lambda m: f"{m // 60}h{m % 60:02d}" if m % 60 else f"{m // 60}h"  # noqa: E731
    for start in range(8 * 60, 18 * 60, 30):
        closed = "X" if 12 * 60 <= start < 13 * 60 else None
        sh.append([f"{fmt(start)} - {fmt(start + 30)}", *[closed] * 5])
    sh.merge_cells("D12:D21")  # mercredi 13h-18h : fermé
    for c in range(1, 7):
        sh.cell(1, c).fill, sh.cell(1, c).font = HEAD_FILL, HEAD_FONT
        sh.column_dimensions[get_column_letter(c)].width = 16
    sh.cell(23, 1, "Grille des créneaux où l'EPS est possible, au pas de 30 min. Case vide = ouvert ; « X » ou "
                   "cellules fusionnées vides = fermé. Les heures de la première colonne donnent la durée des créneaux.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["parse_timetable", "export_solutions", "template_workbook", "TimetableError"]
