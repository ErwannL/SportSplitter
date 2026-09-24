"""Import de la grille Excel et export des plannings."""

from __future__ import annotations

import io
from datetime import time

import pytest
from openpyxl import Workbook, load_workbook

from app.excel_io import (_hex, _text_color, cell_contents, export_solutions, fill_durations, parse_entries,
                          parse_time_label, parse_timetable, template_workbook, to_minutes, TimetableError)
from app.schemas import Solution, TimeRow, Violation
from app.solver import solve

from .factories import grid, level, place, slots, sport, workspace


def book(rows, merges=()) -> bytes:
    wb = Workbook()
    sh = wb.active
    for r in rows:
        sh.append(r)
    for m in merges:
        sh.merge_cells(m)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------- times

def test_to_minutes_and_labels():
    assert to_minutes("8h") == 480 and to_minutes("8h30") == 510 and to_minutes("08:30") == 510
    assert to_minutes("8") == 480 and to_minutes(" 9H05 ") == 545
    assert to_minutes("25h") is None and to_minutes("8h75") is None and to_minutes("matin") is None
    r = parse_time_label("8h - 10h")
    assert (r.label, r.start, r.end) == ("8h – 10h", "8h", "10h")
    assert parse_time_label("8h\n10h").end == "10h"
    assert parse_time_label(time(13, 45)).start == "13h45"
    assert parse_time_label("-").start == "-"


def test_fill_durations():
    rows = [TimeRow(label="8h", start="8h"), TimeRow(label="a", start="9h30", end="10h30"),
            TimeRow(label="10h - 9h", start="10h", end="9h"), TimeRow(label="soir", start="soir")]
    fill_durations(rows)
    assert [r.minutes for r in rows] == [90, 60, 0, 0]
    assert rows[0].label == "8h – 9h30" and rows[1].label == "a" and rows[3].end == ""


def test_parse_entries_still_parses_text():
    e = parse_entries("6eme, 4eme x2; 3e (2)\n2 x Terminale, 6eme,")
    assert [(x.level, x.groups) for x in e] == [("6eme", 2), ("4eme", 2), ("3e", 2), ("Terminale", 2)]


# --------------------------------------------------------------- parse

def test_template_roundtrip():
    tt = parse_timetable(template_workbook(), "modele.xlsx")
    assert tt.file_name == "modele.xlsx"
    assert tt.days == ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi"]
    assert len(tt.rows) == 10 and all(r.minutes == 60 for r in tt.rows)
    assert tt.rows[0].label == "8h – 9h"
    assert all(c.closed for c in tt.cells if c.row == 4)
    wed = [c for c in tt.cells if c.day == 2]
    assert len(wed) == 6 and wed[-1].closed and (wed[-1].row, wed[-1].row_span) == (5, 5)
    assert not any(c.closed for c in tt.cells if c.day == 0 and c.row != 4)


def test_parse_time_objects_offset_header_and_merges():
    data = book([
        ["Planning EPS"],
        [None],
        ["Heures", "Lundi", None, "Mardi"],
        [None],
        [time(8, 0), None, None, "fermé"],
        [time(9, 30), "X", None, None],
        [time(10, 30), "6e x2", None, None],
        ["n'importe quoi", None, None, None],
        [None],
        ["ignored", None, None, None],
    ], merges=["B7:B8", "D6:D7", "D8:E8"])
    tt = parse_timetable(data, "f.xlsx")
    assert tt.days == ["Lundi", "Mardi"]
    assert [r.minutes for r in tt.rows] == [90, 60, 0, 0]
    assert tt.rows[0].label == "8h – 9h30"
    cells = {(c.day, c.row): c for c in tt.cells}
    assert cells[(0, 1)].closed and not cells[(0, 0)].closed
    assert cells[(0, 2)].row_span == 2 and not cells[(0, 2)].closed  # fusion non vide = ouverte
    assert cells[(0, 2)].entries[0].groups == 2
    assert cells[(1, 0)].closed  # « fermé »
    assert cells[(1, 1)].row_span == 2 and cells[(1, 1)].closed  # fusion vide = fermée
    assert cells[(1, 3)].row_span == 1 and not cells[(1, 3)].closed  # fusion horizontale
    assert len(tt.cells) == 6


def test_parse_errors():
    with pytest.raises(TimetableError, match="illisible"):
        parse_timetable(b"not excel")
    with pytest.raises(TimetableError, match="en-tête"):
        parse_timetable(book([[1, 2], [3, 4]]))
    with pytest.raises(TimetableError, match="créneau"):
        parse_timetable(book([["Heures", "Lundi"]]))


# --------------------------------------------------------------- export

def test_colors():
    assert _hex("#abc") == "AABBCC" and _hex("") == "FFFFFF" and _hex("#12345678") == "123456"
    assert _text_color("#ffffff") == "000000" and _text_color("#000") == "FFFFFF"


def _solved(**settings):
    tt = grid(days=2, minutes=[60, 60], closed={(1, 1)}, spans={(0, 0): 2})
    ws = workspace(tt, [level("L", ["s"], cycle=[[120], [60]], groups=2, name="6e")], [sport("s", ["gym"])],
                   [place("gym", slots(tt), capacity=2)], separatePlacesRule="off", **settings)
    ws.places[0].color = "#fff"
    res = solve(ws)
    assert res.status == "ok"
    return ws, res.solutions


def test_cell_contents_filters_segment_and_week():
    ws, sols = _solved()
    sol = sols[0]
    week_a = cell_contents(ws, sol, "Q1", 0)
    assert list(week_a) == [(0, 0), (0, 1)]
    assert week_a[(0, 0)][0] == ("6e ×2 · S · Gym", "#fff")
    week_b = cell_contents(ws, sol, "Q1", 1)
    assert len(week_b) == 1


def test_export_single_fr_multi_week():
    ws, sols = _solved()
    ws.levels.append(level("X", ["s"]))  # absent du plan : ignoré dans le récapitulatif
    wb = load_workbook(io.BytesIO(export_solutions(ws, sols[:1])))
    assert wb.sheetnames == ["Récapitulatif", "Sept - Nov", "Déc - Janv", "Févr - Mars", "Avr - Juin"]
    summ = wb["Récapitulatif"]
    assert [summ.cell(2, c).value for c in range(1, 5)] == ["6e", "Trimestre", "S", "S"]
    assert summ.cell(3, 1).value is None
    q1 = wb["Sept - Nov"]
    assert q1.cell(1, 1).value == "Sept - Nov · Semaine A"
    assert q1.cell(3, 2).value == "6e ×2 · S · Gym"
    assert "B3:B4" in {str(r) for r in q1.merged_cells.ranges}
    assert q1.cell(4, 3).fill.fgColor.rgb.endswith("E2E8F0")  # fermé
    assert q1.cell(7, 1).value == "Sept - Nov · Semaine B"


def test_export_multi_en_with_violations():
    tt = grid(days=1, minutes=[60])
    ws = workspace(tt, [level("L", ["foot"], mode="semestre")], [sport("foot", ["field"])],
                   [place("field", slots(tt), outdoor=True)], maxWinterViolations=5)
    res = solve(ws)
    assert res.status == "relaxed" and len(res.solutions) == 1
    extra = Solution(index=1, plan={"L": {"S1": "foot"}}, assignments=[],
                     violations=[Violation(rule="custom", message="raw message")])
    wb = load_workbook(io.BytesIO(export_solutions(ws, [res.solutions[0], extra], "en")))
    assert wb.sheetnames == ["Solution 1", "Solution 2"]
    s1 = wb["Solution 1"]
    texts = [s1.cell(r, 1).value for r in range(1, 12)]
    assert "Broken rules" in texts
    assert any(isinstance(t, str) and t.startswith("“L” uses an outdoor place (Field) during winter (1st")
               for t in texts)
    assert "Sep - Nov" in texts
    s2 = wb["Solution 2"]
    assert [s2.cell(2, c).value for c in range(1, 5)] == ["L", "Semester", "Foot", None]
    assert "raw message" in [s2.cell(r, 1).value for r in range(1, 8)]


def test_export_without_timetable():
    ws, sols = _solved()
    ws.timetable = None
    with pytest.raises(TimetableError):
        export_solutions(ws, sols)
