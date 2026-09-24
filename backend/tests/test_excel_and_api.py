import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import main
from app.db import Store
from app.excel_io import TimetableError, parse_entries, parse_time_label, parse_timetable, template_workbook
from app.solver import solve
from app.validation import validate

from .factories import level, place, sport, timetable, workspace


def test_parse_entries_groups():
    e = parse_entries("6eme, 4eme x2; 3e (2)\n2 x Terminale, 6eme")
    assert [(x.level, x.groups) for x in e] == [("6eme", 2), ("4eme", 2), ("3e", 2), ("Terminale", 2)]


def test_parse_time_label():
    r = parse_time_label("8h - 10h")
    assert (r.start, r.end) == ("8h", "10h")
    assert parse_time_label("8h\n10h").end == "10h"


def test_template_roundtrip_with_merged_closed_cell():
    tt = parse_timetable(template_workbook(), "modele.xlsx")
    assert tt.days == ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi"]
    assert len(tt.rows) == 5
    wed = [c for c in tt.cells if c.day == 2]
    assert wed[-1].closed and wed[-1].row_span == 2 and wed[-1].row == 3
    assert len(wed) == 4
    lundi_8 = next(c for c in tt.cells if c.day == 0 and c.row == 0)
    assert lundi_8.entries[0].level == "6eme"


def test_parse_rejects_garbage():
    with pytest.raises(TimetableError):
        parse_timetable(b"not excel")
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(TimetableError):
        parse_timetable(buf.getvalue())


def test_validate_flags_missing_config():
    ws = workspace(timetable({(0, 0): [("6e", 1)]}), [level("6e", ["x"])], [sport("x", [])], [])
    codes = {i.code for i in validate(ws)}
    assert "sport_no_place" in codes


@pytest.fixture
def client(tmp_path):
    store = Store(f"sqlite:///{tmp_path}/t.db")
    main.app.dependency_overrides[main.get_store] = lambda: store
    yield TestClient(main.app)
    main.app.dependency_overrides.clear()


def _ws():
    return workspace(
        timetable({(0, 0): [("6e", 1)]}),
        [level("6e", ["bad", "danse"], mode="semestre")],
        [sport("bad", ["gym"]), sport("danse", ["gym"])],
        [place("gym", ["0-0"])],
    )


def test_workspace_persistence(client):
    assert client.get("/api/workspace").json()["levels"] == []
    body = _ws().model_dump(by_alias=True, mode="json")
    assert client.put("/api/workspace", json=body).status_code == 200
    got = client.get("/api/workspace").json()
    assert got["levels"][0]["sportIds"] == ["bad", "danse"]


def test_parse_endpoint(client):
    r = client.post("/api/timetable/parse", files={"file": ("m.xlsx", template_workbook())})
    assert r.status_code == 200 and r.json()["fileName"] == "m.xlsx"
    r = client.post("/api/timetable/parse", files={"file": ("m.xlsx", b"zzz")})
    assert r.status_code == 422


def test_solve_and_export(client):
    ws = _ws()
    body = ws.model_dump(by_alias=True, mode="json")
    res = client.post("/api/solve", json=body).json()
    assert res["status"] == "ok" and res["totalFound"] == 2
    r = client.post("/api/export", json={"workspace": body, "solutions": res["solutions"][:1]})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Récapitulatif"
    r = client.post("/api/export", json={"workspace": body, "solutions": res["solutions"]})
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Solution 1", "Solution 2"]


def test_solve_template_end_to_end():
    tt = parse_timetable(template_workbook())
    names = ["6eme", "5eme", "4eme", "3eme"]
    slots = [c.slot_id for c in tt.cells if not c.closed]
    ws = workspace(
        tt,
        [level(n, ["bad", "co", "danse"], name=n) for n in names],
        [sport("bad", ["gym"]), sport("co", ["parc"], barrette=True), sport("danse", ["salle", "gym"])],
        [place("gym", slots, capacity=2), place("parc", slots, outdoor=True, capacity=2), place("salle", slots)],
    )
    ws.settings.max_solutions = 3
    res = solve(ws)
    # co (barrette) n'est possible que pour la 4eme ; les autres n'ont que 2 sports pour 3 trimestres
    assert res.status in ("ok", "relaxed")
    assert res.solutions
