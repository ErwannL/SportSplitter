"""Tests complémentaires : réglages du solveur, import/export Excel, API, messages."""

import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import main
from app.db import Store
from app.excel_io import TimetableError, _text_color, export_solutions, parse_timetable, template_workbook
from app.messages import render
from app.schemas import Solution, Violation
from app.solver import diagnose, solve
from app.validation import validate

from .factories import level, place, sport, timetable, workspace


def _fast(ws, **settings):
    ws.settings.time_limit = 5
    ws.settings.max_solutions = 20
    for k, v in settings.items():
        setattr(ws.settings, k, v)
    return ws


def _codes(issues):
    return {i.code for i in issues}


# ------------------------------------------------------------------ messages

def test_render_fr_en_and_period_translation():
    p = {"level": "6e", "place": "Parc", "period": "T2"}
    assert "2e trimestre" in render("winter_outdoor", p)
    assert "2nd term" in render("winter_outdoor", p, "en")
    assert "Q9" in render("winter_outdoor", {**p, "period": "Q9"}, "en")
    assert render("no_timetable", {}, "en") == "No timetable imported."


# ---------------------------------------------------------------- validation

def test_validation_codes():
    assert _codes(validate(workspace(None, [], [], []))) == {"no_timetable"}
    ws = workspace(
        timetable({(0, 0): [("6e", 1), ("ghost", 1), ("Ghost", 1)], (0, 1): "closed"}),
        [level("6e", []), level("5e", ["x"]), level("dup", ["x"], name="5E ")],
        [sport("x", ["gym"])],
        [place("gym", {"0-0": []})],
    )
    issues = validate(ws)
    codes = [i.code for i in issues]
    assert codes.count("unknown_level") == 1
    assert {"level_no_sport", "level_unused", "place_never_available", "duplicate_level"} <= set(codes)


# -------------------------------------------------------------------- solver

def _winter_ws():
    # un seul sport, extérieur : répété sur T1, T2, T3 (T2 et T3 touchent l'hiver)
    return workspace(timetable({(0, 0): [("6e", 1)]}), [level("6e", ["foot"])],
                     [sport("foot", ["parc"])], [place("parc", ["0-0"], outdoor=True)])


def test_winter_limit():
    res = solve(_fast(_winter_ws(), max_winter_violations=0))
    assert res.status == "infeasible" and _codes(res.issues) == {"winter_limit"}


def test_winter_soft_relaxed_and_export_violations():
    ws = _fast(_winter_ws(), max_winter_violations=5)
    res = solve(ws)
    assert res.status == "relaxed" and "relaxed" in _codes(res.issues)
    sol = res.solutions[0]
    assert sol.violations and sol.violations[0].rule == "winter_outdoor"
    sol.violations.append(Violation(rule="custom", message="règle brute"))
    for lang in ("fr", "en"):
        wb = load_workbook(io.BytesIO(export_solutions(ws, [sol], lang)))
        values = [c.value for row in wb.worksheets[0].iter_rows() for c in row]
        assert "règle brute" in values
        assert ("Broken rules" if lang == "en" else "Règles non respectées") in values


def test_winter_off_and_hard():
    assert solve(_fast(_winter_ws(), winter_rule="off")).status == "ok"
    res = solve(_fast(_winter_ws(), winter_rule="hard"))
    assert res.status == "infeasible" and "no_solution" in _codes(res.issues)


def test_priority_preferred_maximises_priority():
    ws = workspace(timetable({(0, 0): [("6e", 1)]}), [level("6e", ["a", "b", "c", "d"], mode="semestre")],
                   [sport("a", ["gym"], priority=True), sport("b", ["gym"], priority=True),
                    sport("c", ["gym"], priority=True), sport("d", ["gym"])],
                   [place("gym", ["0-0"])])
    res = solve(_fast(ws))
    assert res.status == "infeasible" and "too_many_priority" in _codes(res.issues)
    res = solve(_fast(ws, priority_required=False))
    assert res.status == "ok" and res.solutions
    assert all("d" not in sol.plan["6e"].values() for sol in res.solutions)


def test_not_enough_sports_when_repeat_disabled():
    ws = workspace(timetable({(0, 0): [("6e", 1)]}), [level("6e", ["a"], mode="semestre"), level("5e", ["a"])],
                   [sport("a", ["gym"])], [place("gym", ["0-0"])])
    res = solve(_fast(ws, allow_repeat=False))
    assert res.status == "infeasible" and "not_enough_sports" in _codes(res.issues)


def test_barrette_min_groups():
    ws = workspace(timetable({(0, 0): [("6e", 2)]}), [level("6e", ["co"], mode="semestre")],
                   [sport("co", ["parc"], barrette=True)], [place("parc", ["0-0"], capacity=2)])
    assert solve(_fast(ws)).status == "ok"
    res = solve(_fast(ws, barrette_min_groups=3))
    assert res.status == "infeasible"
    assert {"barrette_single", "level_blocked"} <= _codes(res.issues)


def test_diagnose_ignores_unknown_levels():
    ws = workspace(timetable({(0, 0): [("ghost", 1)]}), [], [], [])
    assert diagnose(ws) == []


# --------------------------------------------------------------------- excel

def _xlsx(rows, merges=()):
    wb = Workbook()
    sh = wb.active
    for r in rows:
        sh.append(r)
    for m in merges:
        sh.merge_cells(m)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_time_cells_gaps_and_closed_words():
    data = _xlsx([
        [None, 1, 2],                       # ligne numérique : pas un en-tête
        ["Heures", "Lundi", None, "Mardi"],  # colonne vide ignorée
        [None],                              # ligne vide avant les créneaux
        [dt.time(8, 0), "6e x2", None, "X"],
        [dt.time(10, 30), "3 x 5e", None, None],
        [dt.datetime(2024, 1, 1, 14, 0), "4e (2) + 6e", None, None],
    ], merges=["D5:D6"])
    tt = parse_timetable(data)
    assert tt.days == ["Lundi", "Mardi"]
    assert [r.label for r in tt.rows] == ["8h", "10h30", "14h"]
    cells = {(c.day, c.row): c for c in tt.cells}
    assert cells[1, 0].closed and cells[1, 1].closed and cells[1, 1].row_span == 2
    assert [(e.level, e.groups) for e in cells[0, 1].entries] == [("5e", 3)]
    assert [(e.level, e.groups) for e in cells[0, 2].entries] == [("4e", 2), ("6e", 1)]


def test_parse_errors():
    with pytest.raises(TimetableError, match="en-tête"):
        parse_timetable(_xlsx([["Heures"], ["8h"]]))
    with pytest.raises(TimetableError, match="créneau"):
        parse_timetable(_xlsx([["Heures", "Lundi"]]))


def test_text_color_and_short_hex():
    assert _text_color("#fff") == "000000"
    assert _text_color("#123") == "FFFFFF"


def test_export_merged_closed_and_unplanned_level():
    tt = parse_timetable(template_workbook())
    slot = next(c for c in tt.cells if c.entries and c.entries[0].level == "6eme").slot_id
    gym = place("gym", [slot])
    gym.color = "#ff0"
    ws = workspace(tt, [level("6eme", ["bad"]), level("autre", ["bad"])], [sport("bad", ["gym"])], [gym])
    sol = Solution(index=0, plan={"6eme": {"T1": "bad"}}, assignments=[
        {"slotId": slot, "levelId": "6eme", "period": "T1", "sportId": "bad",
         "placements": [{"placeId": "gym", "groups": 2}]}])
    for lang in ("fr", "en"):
        wb = load_workbook(io.BytesIO(export_solutions(ws, [sol, sol.model_copy(update={"index": 1})], lang)))
        assert len(wb.sheetnames) == 2
        wb = load_workbook(io.BytesIO(export_solutions(ws, [sol], lang)))
        assert wb.sheetnames[0] == ("Summary" if lang == "en" else "Récapitulatif")
    with pytest.raises(TimetableError):
        export_solutions(workspace(None, [], [], []), [sol])


# ----------------------------------------------------------------------- api

@pytest.fixture
def client(tmp_path):
    store = Store(f"sqlite:///{tmp_path}/t.db")
    main.app.dependency_overrides[main.get_store] = lambda: store
    yield TestClient(main.app)
    main.app.dependency_overrides.clear()


def _body():
    ws = workspace(timetable({(0, 0): [("6e", 1)]}), [level("6e", ["bad"], mode="semestre")],
                   [sport("bad", ["gym"])], [place("gym", ["0-0"])])
    return ws.model_dump(by_alias=True, mode="json")


def test_get_store_default(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    main.get_store.cache_clear()
    try:
        assert isinstance(main.get_store(), Store)
    finally:
        main.get_store.cache_clear()


def test_me_and_roles(client, monkeypatch):
    assert client.get("/api/health").json() == {"status": "ok"}
    assert "edit_rules" in client.get("/api/me").json()["permissions"]
    monkeypatch.setenv("SPORTSPLITTER_ROLE", "teacher")
    me = client.get("/api/me").json()
    assert me == {"role": "teacher", "permissions": ["edit_workspace"]}
    body = _body()
    assert client.put("/api/workspace", json=body).status_code == 200
    assert client.put("/api/workspace", json=body).status_code == 200  # mise à jour (upsert)
    body["settings"]["maxSolutions"] = 3
    assert client.put("/api/workspace", json=body).status_code == 403


def test_template_validate_and_export_errors(client):
    r = client.get("/api/timetable/template")
    assert r.status_code == 200 and parse_timetable(r.content).days[0] == "Lundi"
    body = _body()
    assert client.post("/api/validate", json=body).json() == []
    assert client.post("/api/export", json={"workspace": body, "solutions": []}).status_code == 422
    sols = client.post("/api/solve", json=body).json()["solutions"]
    r = client.post("/api/export", json={"workspace": {**body, "timetable": None}, "solutions": sols})
    assert r.status_code == 422
    r = client.post("/api/export", json={"workspace": body, "solutions": sols, "lang": "en"})
    assert r.status_code == 200 and load_workbook(io.BytesIO(r.content)).sheetnames[0] == "Summary"
