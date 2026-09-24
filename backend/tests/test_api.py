"""API HTTP et persistance."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.db import Store
from app.excel_io import parse_timetable, template_workbook

from .conftest import login
from .factories import grid, level, place, slots, sport, workspace

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _body(max_solutions=5):
    tt = grid(days=2, minutes=[60])
    ws = workspace(tt, [level("6e", ["s"], mode="semestre")], [sport("s", ["gym"])], [place("gym", slots(tt))],
                   maxSolutions=max_solutions)
    return ws.model_dump(by_alias=True, mode="json")


def test_db_insert_then_update(tmp_path):
    store = Store(f"sqlite:///{tmp_path}/db.sqlite")
    assert store.load().levels == []
    ws = workspace(grid(), [level("A", ["s"])], [], [])
    store.save(ws)
    ws.levels[0].name = "B"
    store.save(ws)
    assert store.load().levels[0].name == "B"


def test_me_roles_and_workspace(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    assert client.get("/api/me").json() == {"sub": "4821", "email": "4821@orqea.dev", "name": "User 4821",
                                            "role": "admin", "permissions": ["edit_workspace", "edit_rules"]}
    body = _body()
    body["settings"]["maxSolutions"] = 7
    assert client.put("/api/workspace", json=body).status_code == 200  # admin : règles modifiables
    login(client, role="user")
    assert client.get("/api/me").json()["permissions"] == ["edit_workspace"]
    body["levels"][0]["name"] = "Sixième"
    assert client.put("/api/workspace", json=body).status_code == 200  # contenu modifiable
    body["settings"]["maxSolutions"] = 3
    r = client.put("/api/workspace", json=body)
    assert r.status_code == 403
    saved = client.get("/api/workspace").json()
    assert saved["levels"][0]["name"] == "Sixième" and saved["settings"]["maxSolutions"] == 7


def test_parse_and_template(client):
    r = client.get("/api/timetable/template")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert parse_timetable(r.content).days[0] == "Lundi"
    r = client.get("/api/timetable/template?saturday=true")
    assert parse_timetable(r.content).days[-1] == "Samedi"
    r = client.post("/api/timetable/parse", files={"file": ("m.xlsx", template_workbook(), XLSX)})
    assert r.status_code == 200 and r.json()["fileName"] == "m.xlsx" and r.json()["rows"][0]["minutes"] == 30
    r = client.post("/api/timetable/parse", files={"file": ("bad.xlsx", b"nope", XLSX)})
    assert r.status_code == 422 and "illisible" in r.json()["detail"]


def test_validate_solve_export(client):
    body = _body()
    assert client.post("/api/validate", json=body).json() == []
    bad = {**body, "levels": []}
    assert [i["code"] for i in client.post("/api/validate", json=bad).json()] == ["no_level"]
    res = client.post("/api/solve", json=body).json()
    assert res["status"] == "ok" and res["totalFound"] == 2 and res["solutions"][0]["weeks"] == 1
    sols = res["solutions"]
    assert client.post("/api/export", json={"workspace": body, "solutions": []}).status_code == 422
    r = client.post("/api/export", json={"workspace": {**body, "timetable": None}, "solutions": sols})
    assert r.status_code == 422
    r = client.post("/api/export", json={"workspace": body, "solutions": sols[:1], "lang": "en"})
    assert r.status_code == 200 and 'filename="planning.xlsx"' in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).sheetnames[0] == "Summary"
    r = client.post("/api/export", json={"workspace": body, "solutions": sols})
    assert 'filename="plannings.xlsx"' in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).sheetnames == [f"Solution {i}" for i in range(1, 3)]
